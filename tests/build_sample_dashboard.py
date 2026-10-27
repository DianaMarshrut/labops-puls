"""Run the puls algorithm (steps 4-5) on a sample xlsx and fill the template.

Reads the deals export, computes 6 metrics (current 4 weeks vs previous
4 weeks), and writes a filled dashboard HTML. This mirrors what the AI does
inside a chat; it exists only to test the template end to end.
"""

from __future__ import annotations

import json
import logging
import re
import sys
from collections.abc import Callable
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

logger = logging.getLogger(__name__)
NO_DATA = "нет данных"
REPO = Path(__file__).resolve().parent.parent
TEMPLATE = REPO / "plugins/puls/skills/puls/templates/dashboard.html"


@dataclass
class Deal:
    """One row of the deals export."""

    stage: str
    amount: float
    manager: str
    created: date
    closed: date | None
    changed: date


@dataclass
class Kpi:
    """One KPI card."""

    label: str
    value: str
    delta_pct: float | None
    better: str = "up"
    delta_unit: str = "%"
    compare_to: str = "прошлым 4 неделям"
    note: str = ""


@dataclass
class Dashboard:
    """Template payload."""

    meta: dict[str, str]
    kpis: list[Kpi]
    charts: list[dict] = field(default_factory=list)
    tables: list[dict] = field(default_factory=list)
    attention: list[str] = field(default_factory=list)


def _d(v: object) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def read_deals(path: Path) -> list[Deal]:
    """Read deals from the first sheet; header row maps columns by name."""
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = ws.iter_rows(values_only=True)
    head = {str(h).strip(): i for i, h in enumerate(next(rows)) if h is not None}
    out: list[Deal] = []
    for r in rows:
        if r[head["ID"]] is None:
            continue
        out.append(Deal(
            stage=str(r[head["Стадия сделки"]]),
            amount=float(r[head["Сумма"]] or 0),
            manager=str(r[head["Ответственный"]]),
            created=_d(r[head["Дата создания"]]),
            closed=_d(r[head["Дата закрытия"]]),
            changed=_d(r[head["Дата изменения"]]),
        ))
    return out


def read_all_pages(fetch: Callable[[int], dict]) -> list[dict]:
    """Read every page of a Bitrix24 list answer, as SKILL.md step 4 tells the AI to do.

    Args:
        fetch: Opens the saved link with the given `start` and returns the JSON answer.

    Returns:
        All records. Stops only when the answer has no `next`, never by dates.
    """
    rows: list[dict] = []
    start = 0
    while True:
        answer = fetch(start)
        rows += answer["result"]
        if "next" not in answer:
            return rows
        start = answer["next"]


def money(v: float) -> str:
    """Format rubles with thin thousands separators, no kopecks."""
    return f"{round(v):,}".replace(",", " ") + " ₽"


def pct(cur: float | None, prev: float | None) -> float | None:
    """Percent change; None when either value is missing or the previous one is zero."""
    if cur is None or prev is None or prev == 0:
        return None
    return round((cur - prev) / prev * 100, 1)


def compute(deals: list[Deal], today: date, data_from: date) -> Dashboard:
    """Compute metrics for the last 28 days vs the 28 days before.

    Args:
        deals: Rows of the export.
        today: Export date (end of the current period).
        data_from: First day the export covers, as stated by the user. Periods
            starting before it are "no data", never zero.
    """
    start = today - timedelta(days=27)
    pstart = start - timedelta(days=28)

    def won(a: date, b: date) -> list[Deal]:
        return [d for d in deals if d.stage == "Успешно" and d.closed and a <= d.closed <= b]

    def closed(a: date, b: date) -> list[Deal]:
        return [d for d in deals if d.closed and a <= d.closed <= b]

    def new(a: date, b: date) -> list[Deal]:
        return [d for d in deals if a <= d.created <= b]

    pend = start - timedelta(days=1)
    has_cur = data_from <= start
    has_prev = data_from <= pstart
    w, pw = won(start, today), won(pstart, pend)
    c, pc = closed(start, today), closed(pstart, pend)
    n, pn = new(start, today), new(pstart, pend)
    rev = sum(d.amount for d in w) if has_cur else None
    prev_rev = sum(d.amount for d in pw) if has_prev else None
    n_cur = len(n) if has_cur else None
    n_prev = len(pn) if has_prev else None
    conv = len(w) / len(c) * 100 if has_cur and c else None
    pconv = len(pw) / len(pc) * 100 if has_prev and pc else None
    avg = rev / len(w) if rev is not None and w else None
    pavg = prev_rev / len(pw) if prev_rev is not None and pw else None
    open_deals = [d for d in deals if not d.closed]
    stale = [d for d in open_deals if (today - d.changed).days > 14]

    def show(v: float | None, f: Callable[[float], str]) -> str:
        return NO_DATA if v is None else f(v)

    kpis = [
        Kpi("Выручка (успешные сделки)", show(rev, money), pct(rev, prev_rev)),
        Kpi("Новые сделки", show(n_cur, lambda v: str(int(v))), pct(n_cur, n_prev)),
        Kpi("Конверсия в успех", show(conv, lambda v: f"{v:.0f}%"),
            round(conv - pconv, 1) if conv is not None and pconv is not None else None,
            delta_unit=" п.п."),
        Kpi("Средний чек", show(avg, money), pct(avg, pavg)),
        Kpi("Сделок в работе", str(len(open_deals)), None, note="Сейчас, без сравнения"),
        Kpi("Без движения > 14 дней", str(len(stale)), None, better="down",
            note="Открытые сделки, которые давно не менялись"),
    ]
    weeks = []
    for k in range(4):
        a = start + timedelta(days=7 * k)
        b = a + timedelta(days=6)
        val = round(sum(d.amount for d in won(a, b))) if data_from <= a else None
        weeks.append({"label": f"{a:%d.%m}–{b:%d.%m}", "value": val})
    by_mgr: dict[str, list[Deal]] = {}
    for d in (w if has_cur else []):
        by_mgr.setdefault(d.manager, []).append(d)
    rows = sorted(by_mgr.items(), key=lambda kv: -sum(x.amount for x in kv[1]))
    stages = ["Новая", "Переговоры", "Счёт выставлен"]
    attention = []
    if stale:
        attention.append(f"{len(stale)} открытых сделок не менялись больше 14 дней.")
    dr = pct(rev, prev_rev)  # None-safe
    if dr is not None and abs(dr) >= 10:
        word = "выросла" if dr > 0 else "упала"
        attention.append(f"Выручка {word} на {abs(dr):.0f}% к прошлым 4 неделям.")
    return Dashboard(
        meta={
            "company": "ООО «Пример» — отдел продаж",
            "period": f"Период: {start:%d.%m}–{today:%d.%m.%Y} (сравнение с {pstart:%d.%m}–{pend:%d.%m})",
            "built_at": f"{today:%d.%m.%Y} 15:00 (UTC+5)",
            "source": "sample-deals.xlsx (выгрузка сделок, тестовые данные)",
        },
        kpis=kpis,
        charts=[
            {"title": "Выручка по неделям", "type": "bar", "unit": "₽", "points": weeks},
            {"title": "Сделки в работе по стадиям", "type": "bar", "unit": "шт",
             "points": [{"label": s, "value": sum(1 for d in open_deals if d.stage == s)}
                        for s in stages]},
        ],
        tables=[{
            "title": "Выручка по менеджерам",
            "columns": ["Менеджер", "Успешных сделок", "Выручка"],
            "num_cols": [1, 2],
            "empty_note": "" if has_cur else "нет данных: выгрузка не охватывает весь период",
            "rows": [[m, str(len(ds)), money(sum(x.amount for x in ds))] for m, ds in rows],
        }],
        attention=attention[:3],
    )


WEBHOOK_KEY = re.compile(r"/rest/\d+/[^/\s\"?#]+")


def check_no_secrets(html: str) -> None:
    """Raise if a webhook link with a key is anywhere in the finished page (SKILL.md, step 2).

    Args:
        html: The whole filled dashboard.

    Raises:
        ValueError: When a ".../rest/<id>/<key>" link is found.
    """
    if WEBHOOK_KEY.search(html):
        raise ValueError("webhook link with a key found in the dashboard")


def fill(template: Path, dash: Dashboard, out: Path) -> Path:
    """Replace the JSON payload inside the template, check it for secrets and save."""
    html = template.read_text(encoding="utf-8")
    a = html.index('<script id="puls-data" type="application/json">')
    a = html.index(">", a) + 1
    b = html.index("</script>", a)
    payload = json.dumps(asdict(dash), ensure_ascii=False, indent=1).replace("<", "\\u003c")
    page = html[:a] + "\n" + payload + "\n" + html[b:]
    check_no_secrets(page)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(page, encoding="utf-8")
    logger.info("Saved %s", out)
    return out


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else REPO / "tests/out/sample-deals.xlsx"
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else REPO / "tests/out/pulse-sample.html"
    today = date(2026, 10, 27)
    # The sample export covers deals created in the last 60 days (see make_sample_xlsx.py).
    fill(TEMPLATE, compute(read_deals(src), today, today - timedelta(days=60)), dst)
